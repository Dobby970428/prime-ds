"""포터블 배포본 자체 검증 — 복사가 끝난 뒤 배포본의 런타임으로 실행합니다.

«돌아가는지» 가 아니라 «논문과 같은 답을 내는지» 까지 봅니다.
하나라도 실패하면 0 이 아닌 코드로 끝나고, 빌드가 중단됩니다.

검사 항목
  1  필수 패키지 import 와 버전
  2  모델 피클에 기록된 scikit-learn 버전 == 설치된 버전   ← 조용한 오작동 방지
  3  Layer C 모델 적재 + 알려진 구조 2개 예측
  4  Layer A 실측 DB 열기 + 행 수
  5  구조 경보 계층 — NDMA 가 N-nitroso 경보에 걸리는지
  6  Flask 앱 생성 + 등록된 모든 GET 화면이 500 을 내지 않는지
  7  보고서 출력 경로 (python-docx · fpdf2 · openpyxl) 적재
"""
from __future__ import annotations

import re
import sqlite3
import sys
import traceback
from pathlib import Path

APP = Path(__file__).resolve().parent
if str(APP) not in sys.path:
    sys.path.insert(0, str(APP))

OK, NG = [], []


def check(name):
    def deco(fn):
        try:
            detail = fn()
            OK.append((name, detail or ""))
        except Exception as exc:
            NG.append((name, f"{type(exc).__name__}: {exc}"))
            traceback.print_exc()
        return fn
    return deco


# ── 1 ────────────────────────────────────────────────────────────────
@check("필수 패키지")
def _():
    from importlib.metadata import version
    import flask, numpy, pandas, sklearn, rdkit, PIL, requests   # noqa: F401
    import docx, fpdf, openpyxl, rdchiral                        # noqa: F401
    from rdkit import rdBase
    # flask.__version__ 은 Flask 3.2 에서 사라집니다. metadata 로 읽습니다.
    return (f"python {sys.version.split()[0]} · rdkit {rdBase.rdkitVersion} · "
            f"sklearn {sklearn.__version__} · numpy {numpy.__version__} · "
            f"pandas {pandas.__version__} · flask {version('flask')} · "
            f"pillow {version('pillow')} · fpdf2 {version('fpdf2')}")


# ── 2 ────────────────────────────────────────────────────────────────
@check("모델 피클 ↔ scikit-learn 버전")
def _():
    import sklearn
    raw = (APP / "models" / "ames_ecfp_rf.pkl").read_bytes()
    m = re.search(rb"_sklearn_version\x94\x8c.(\d+\.\d+\.\d+)", raw)
    if not m:
        m = re.search(rb"_sklearn_version.{0,8}?(\d+\.\d+\.\d+)", raw, re.S)
    if not m:
        raise RuntimeError("피클에서 _sklearn_version 을 찾지 못했습니다")
    pickled = m.group(1).decode()
    if pickled != sklearn.__version__:
        raise RuntimeError(
            f"모델은 scikit-learn {pickled} 로 저장되었는데 설치된 것은 "
            f"{sklearn.__version__} 입니다. 배포\\requirements-dist.txt 의 "
            f"scikit-learn 고정 버전을 {pickled} 으로 맞추십시오."
        )
    return f"양쪽 모두 {pickled}"


# ── 3 ────────────────────────────────────────────────────────────────
@check("Layer C 모델 적재 · 예측")
def _():
    # 손으로 지문을 다시 만들지 않고 플랫폼이 실제로 쓰는 경로를 그대로 탄다.
    import p63_deepchem_predictor as p63

    st = p63.model_status()
    if not st.get("ready"):
        raise RuntimeError(f"모델이 준비되지 않았습니다 — {st.get('error')}")
    if not st.get("endpoint_is_ames"):
        # Tox21 SR-p53 대체 모델이 잡히면 판정이 통째로 달라진다.
        raise RuntimeError(
            f"엔드포인트가 Ames 가 아닙니다 ({st.get('endpoint')}). "
            f"models/ames_ecfp_rf.pkl 이 빠졌거나 환경변수 "
            f"PRIME_LAYER_C_MODEL 이 설정되어 있습니다."
        )
    rows = p63.predict_advanced_properties(["CN(C)N=O", "CC(=O)O"])
    out = []
    for name, smi in (("NDMA", "CN(C)N=O"), ("아세트산", "CC(=O)O")):
        row = rows.get(smi) or {}
        p = row.get("Ames_Mutagenicity_Prob")
        if p is None:
            raise RuntimeError(f"{name}: 예측값이 비어 있습니다 ({row.get('Prediction_Class')})")
        out.append(f"{name} P={float(p):.3f}")
    return " · ".join(out)


# ── 4 ────────────────────────────────────────────────────────────────
@check("Layer A 실측 DB")
def _():
    db = APP / "cache" / "m7_empirical.db"
    con = sqlite3.connect(f"file:{db}?mode=ro", uri=True)
    try:
        tables = [r[0] for r in con.execute(
            "SELECT name FROM sqlite_master WHERE type='table'")]
        if not tables:
            raise RuntimeError("테이블이 없습니다")
        counts = []
        for t in tables:
            counts.append(f"{t} {con.execute(f'SELECT COUNT(*) FROM \"{t}\"').fetchone()[0]}행")
        return " · ".join(counts)
    finally:
        con.close()


# ── 5 ────────────────────────────────────────────────────────────────
@check("구조 경보 계층")
def _():
    from rdkit import Chem
    import p62_toxicity_alerts as p62

    alerts, _idx = p62.scan_ashby_tennant_alerts(Chem.MolFromSmiles("CN(C)N=O"))
    if not alerts:
        raise RuntimeError("NDMA 가 어떤 경보에도 걸리지 않았습니다 — 경보표가 빠졌습니다")
    clean, _i2 = p62.scan_ashby_tennant_alerts(Chem.MolFromSmiles("CC(=O)O"))
    if clean:
        raise RuntimeError(f"아세트산에 경보가 걸렸습니다: {clean}")
    n = len(getattr(p62, "ASHBY_TENNANT_ALERTS", []) or [])
    return f"NDMA → {alerts[0]} · 아세트산 → 없음 · 경보표 {n}종"


# ── 6 ────────────────────────────────────────────────────────────────
@check("웹 화면")
def _():
    from webapp import create_app
    app = create_app()
    app.config["TESTING"] = True
    client = app.test_client()
    checked, bad = 0, []
    for rule in app.url_map.iter_rules():
        if "GET" not in rule.methods or rule.arguments:
            continue
        if rule.rule.startswith("/static"):
            continue
        resp = client.get(rule.rule)
        checked += 1
        if resp.status_code >= 500:
            bad.append(f"{rule.rule} → {resp.status_code}")
    if bad:
        raise RuntimeError("500 오류: " + ", ".join(bad))
    return f"GET 화면 {checked}개 모두 정상"


# ── 7 ────────────────────────────────────────────────────────────────
@check("보고서 출력")
def _():
    import io
    from docx import Document
    from fpdf import FPDF
    from openpyxl import Workbook

    d = Document(); d.add_paragraph("PRIME-DS"); b1 = io.BytesIO(); d.save(b1)
    pdf = FPDF(); pdf.add_page(); pdf.set_font("Helvetica", size=11)
    pdf.cell(40, 10, "PRIME-DS"); raw = pdf.output()
    wb = Workbook(); wb.active["A1"] = "PRIME-DS"; b2 = io.BytesIO(); wb.save(b2)
    return f"docx {b1.tell()}B · pdf {len(bytes(raw))}B · xlsx {b2.tell()}B"


# ── 결과 ──────────────────────────────────────────────────────────────
print()
for name, detail in OK:
    print(f"       [통과] {name:28s} {detail}")
for name, detail in NG:
    print(f"       [실패] {name:28s} {detail}")
print()
print(f"       통과 {len(OK)} · 실패 {len(NG)}")

# 검색 데이터 추가팩 안내 (실패로 치지 않음)
idx = APP / "drug_index.db"
if not idx.exists():
    print("       [안내] drug_index.db 없음 — Search 모듈은 비어 있는 상태로 동작합니다")

sys.exit(1 if NG else 0)
